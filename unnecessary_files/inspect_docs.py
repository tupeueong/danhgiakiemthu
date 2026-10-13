import sys
import docx

sys.stdout.reconfigure(encoding='utf-8')

print("Loading Doc 5...")
doc5 = docx.Document(r'C:\Users\dang khoa\Desktop\danhgia_kiemthu\Báo_cáo_Nhóm_5_Testing.docx')
print(f"Doc 5 Paragraphs: {len(doc5.paragraphs)}, Tables: {len(doc5.tables)}")

print("\n--- Doc 5 Headings / Outline ---")
for i, p in enumerate(doc5.paragraphs):
    t = p.text.strip()
    if not t:
        continue
    if p.style.name.startswith('Heading') or any(t.startswith(prefix) for prefix in ['Chương', 'PHỤ LỤC', 'KẾT LUẬN', 'LỜI CẢM ƠN', '1.', '2.', '3.', '4.', '5.', '6.', '7.']):
        if len(t) < 80:
            print(f"[{p.style.name}] {t}")

print("\nLoading Doc 9...")
doc9 = docx.Document(r'C:\Users\dang khoa\Desktop\danhgia_kiemthu\Báo_cáo_Nhóm_9_Testing.docx')
print(f"Doc 9 Paragraphs: {len(doc9.paragraphs)}, Tables: {len(doc9.tables)}")

print("\n--- Doc 9 Headings / Outline ---")
for i, p in enumerate(doc9.paragraphs):
    t = p.text.strip()
    if not t:
        continue
    if p.style.name.startswith('Heading') or any(t.startswith(prefix) for prefix in ['Chương', 'PHỤ LỤC', 'KẾT LUẬN', 'LỜI CẢM ƠN', '1.', '2.', '3.', '4.', '5.', '6.', '7.']):
        if len(t) < 80:
            print(f"[{p.style.name}] {t}")
